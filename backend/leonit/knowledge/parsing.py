"""Разбор загруженных файлов в текст и нарезка на фрагменты.

Форматы: txt, md, html, pdf, docx, csv, xlsx, json. Формат определяется по
расширению, при его отсутствии — по content-type. Всё чужое отклоняется с
понятной ошибкой, а не «падает внутри»: библиотеки разбора бросают самые разные
исключения, здесь они сводятся к ``UnsupportedFormatError``/``ParseError``.
"""

from __future__ import annotations

import csv
import io
import json
import re
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import PurePosixPath
from typing import Any

from leonit.core.errors import ValidationFailedError

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_TEXT_CHARS = 300_000
CHUNK_CHARS = 1200
MIN_CHUNK_CHARS = 200

SUPPORTED_EXTENSIONS: tuple[str, ...] = (
    ".txt",
    ".md",
    ".markdown",
    ".html",
    ".htm",
    ".pdf",
    ".docx",
    ".csv",
    ".xlsx",
    ".json",
)

_CONTENT_TYPES: dict[str, str] = {
    "text/plain": ".txt",
    "text/markdown": ".md",
    "text/html": ".html",
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "text/csv": ".csv",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
    "application/json": ".json",
}

_BLOCK_TAGS = frozenset(
    {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "section", "article"}
)
_SKIP_TAGS = frozenset({"script", "style", "noscript", "svg", "head"})
_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]")
_SPACES_RE = re.compile(r"[ \t ]+")
_BLANK_LINES_RE = re.compile(r"\n{3,}")
_SENTENCE_RE = re.compile(r"(?<=[.!?…])\s+")


class UnsupportedFormatError(ValidationFailedError):
    pass


class ParseError(ValidationFailedError):
    pass


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    text: str
    content_type: str
    extension: str


def normalized_extension(filename: str | None, content_type: str | None) -> str:
    """Расширение в нижнем регистре по имени файла, иначе по content-type."""
    suffix = PurePosixPath(filename or "").suffix.lower()
    if suffix in SUPPORTED_EXTENSIONS:
        return suffix
    base_type = (content_type or "").split(";", 1)[0].strip().lower()
    mapped = _CONTENT_TYPES.get(base_type)
    if mapped:
        return mapped
    if suffix:
        raise UnsupportedFormatError(
            f"Формат {suffix} не поддерживается. Подходят: {', '.join(SUPPORTED_EXTENSIONS)}"
        )
    raise UnsupportedFormatError(
        "Не удалось определить формат файла. Подходят: " + ", ".join(SUPPORTED_EXTENSIONS)
    )


def clean_text(text: str, *, limit: int = MAX_TEXT_CHARS) -> str:
    """Убрать управляющие символы, схлопнуть пробелы, оставить абзацы."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _CONTROL_RE.sub("", text)
    lines = [_SPACES_RE.sub(" ", line).strip() for line in text.split("\n")]
    text = "\n".join(lines).strip()
    text = _BLANK_LINES_RE.sub("\n\n", text)
    return text[:limit]


def extract_text(filename: str | None, content_type: str | None, data: bytes) -> ParsedDocument:
    if len(data) > MAX_FILE_BYTES:
        raise ValidationFailedError(
            f"Файл больше {MAX_FILE_BYTES // (1024 * 1024)} МБ; разбейте его на части"
        )
    extension = normalized_extension(filename, content_type)
    try:
        if extension in (".txt", ".md", ".markdown"):
            text = _decode(data)
        elif extension in (".html", ".htm"):
            text = _html_to_text(_decode(data))
        elif extension == ".pdf":
            text = _pdf_to_text(data)
        elif extension == ".docx":
            text = _docx_to_text(data)
        elif extension == ".csv":
            text = _csv_to_text(_decode(data))
        elif extension == ".xlsx":
            text = _xlsx_to_text(data)
        elif extension == ".json":
            text = _json_to_text(_decode(data))
        else:  # pragma: no cover - normalized_extension уже отсеяла лишнее
            raise UnsupportedFormatError(f"Формат {extension} не поддерживается")
    except ValidationFailedError:
        raise
    except Exception as error:
        raise ParseError(
            f"Не удалось прочитать файл как {extension}: {type(error).__name__}"
        ) from error
    text = clean_text(text)
    if not text:
        raise ParseError("В файле не нашлось текста: сканы и картинки пока не распознаются")
    mapped_type = next((k for k, v in _CONTENT_TYPES.items() if v == extension), "text/plain")
    return ParsedDocument(text=text, content_type=mapped_type, extension=extension)


def chunk_text(text: str, *, size: int = CHUNK_CHARS) -> list[str]:
    """Нарезать текст на фрагменты по абзацам; длинные абзацы — по предложениям.

    Фрагменты не пересекаются и не длиннее ``size`` (кроме одного предложения
    длиннее лимита — его режем по символам); соседние короткие абзацы
    склеиваются, чтобы в поиске не всплывали обрывки из одной строки.
    """
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    pieces: list[str] = []
    for paragraph in paragraphs:
        if len(paragraph) <= size:
            pieces.append(paragraph)
            continue
        current = ""
        for sentence in _SENTENCE_RE.split(paragraph):
            if not sentence:
                continue
            if current and len(current) + len(sentence) + 1 > size:
                pieces.append(current)
                current = sentence
            else:
                current = f"{current} {sentence}".strip()
            while len(current) > size:
                pieces.append(current[:size])
                current = current[size:]
        if current:
            pieces.append(current)
    chunks: list[str] = []
    for piece in pieces:
        if chunks and len(chunks[-1]) + len(piece) + 2 <= size:
            chunks[-1] = f"{chunks[-1]}\n\n{piece}"
        else:
            chunks.append(piece)
    return chunks


# ------------------------------------------------------------------ форматы


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("cp1251", errors="replace")


class _TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in _SKIP_TAGS:
            self._skip += 1
        elif tag in _BLOCK_TAGS:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in _SKIP_TAGS and self._skip:
            self._skip -= 1
        elif tag in _BLOCK_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if not self._skip:
            self.parts.append(data)


def _html_to_text(html: str) -> str:
    parser = _TextExtractor()
    parser.feed(html)
    parser.close()
    return "".join(parser.parts)


def _pdf_to_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as error:
            raise ParseError("PDF защищён паролем") from error
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(page.strip() for page in pages if page.strip())


def _docx_to_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
        parts.append(f"# {text}" if style.startswith("heading") else text)
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append("\n".join(rows))
    return "\n\n".join(parts)


def _csv_to_text(text: str) -> str:
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [
        " | ".join(cell.strip() for cell in row if cell is not None)
        for row in csv.reader(io.StringIO(text), dialect)
        if any(cell.strip() for cell in row)
    ]
    return "\n".join(rows)


def _xlsx_to_text(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if value is None else str(value).strip() for value in row]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"# Лист: {sheet.title}\n" + "\n".join(rows))
    workbook.close()
    return "\n\n".join(parts)


def _json_to_text(text: str) -> str:
    value = json.loads(text)
    lines: list[str] = []

    def walk(item: Any, prefix: str, depth: int) -> None:
        if depth > 6:
            lines.append(f"{prefix}: {json.dumps(item, ensure_ascii=False)[:500]}")
            return
        if isinstance(item, dict):
            for key, child in item.items():
                walk(child, f"{prefix}.{key}" if prefix else str(key), depth + 1)
        elif isinstance(item, list):
            for index, child in enumerate(item):
                walk(child, f"{prefix}[{index}]", depth + 1)
        else:
            lines.append(f"{prefix}: {item}")

    walk(value, "", 0)
    return "\n".join(lines)
