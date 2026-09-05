"""База знаний: разбор форматов, нарезка, поиск, права, ассистент и быстрый черновик вакансии."""

from __future__ import annotations

import io
import json
import uuid

import pytest
from httpx import AsyncClient

from leonit.knowledge import parsing, search
from leonit.knowledge.parsing import ParseError, UnsupportedFormatError, chunk_text, extract_text
from tests.helpers import bearer, create_invite, invite_token_from_url, register
from tests.test_assistant import _actions, _send, _thread

COMPANY_TEXT = """# О компании

Napoleon IT — разработчик программного обеспечения и технологический партнёр
ритейлеров. Компания основана в 2011 году в Челябинске, офисы в Москве и
Санкт-Петербурге.

# Стажировка

NapTalentBootcamp — стажировка для студентов и выпускников. Направления:
бэкенд на Python и FastAPI, наука о данных, тестирование, аналитика.
Гибкий график и наставничество от опытных специалистов.

# Найм

Этапы отбора: отклик на Хантфлоу, разговор с менеджером по персоналу,
интервью по проверке навыков, финальное интервью с руководителем, оффер.
"""


async def _manager(client: AsyncClient, owner: str) -> str:
    invite = await create_invite(client, owner, role="hiring_manager", vacancy_scope=[])
    _, token = await register(client, invite_token=invite_token_from_url(invite["url"]))
    return token


def _docx_bytes() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Ценности", level=1)
    document.add_paragraph("Learn IT — развиваемся, опережая время.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Офис"
    table.rows[0].cells[1].text = "Челябинск, проспект Ленина, 60А"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Льготы"
    sheet.append(["Льгота", "Условие"])
    sheet.append(["ДМС", "после испытательного срока"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _pdf_bytes(text: str = "Napoleon IT knowledge base") -> bytes:
    """Минимальный PDF с одной строкой текста, без внешних библиотек."""
    content = f"BT /F1 18 Tf 40 700 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(out)


# ------------------------------------------------------------------ разбор


def test_extract_text_supports_every_format() -> None:
    assert "О компании" in extract_text("about.md", None, COMPANY_TEXT.encode("utf-8")).text
    assert "привет" in extract_text("win.txt", None, "привет".encode("cp1251")).text
    html = (
        "<html><head><style>p{}</style><script>x()</script></head>"
        "<body><h1>Стек</h1><p>Python</p></body></html>"
    ).encode()
    parsed = extract_text("stack.html", "text/html", html)
    assert "Стек" in parsed.text and "Python" in parsed.text and "x()" not in parsed.text
    csv_text = extract_text("benefits.csv", None, "Льгота;Условие\nДМС;с первого дня\n".encode())
    assert csv_text.text.splitlines()[1] == "ДМС | с первого дня"
    payload = json.dumps({"офисы": ["Москва", "Челябинск"]}).encode()
    as_json = extract_text("meta.json", None, payload)
    assert "офисы[0]: Москва" in as_json.text
    docx = extract_text("values.docx", None, _docx_bytes())
    assert "# Ценности" in docx.text and "Офис | Челябинск, проспект Ленина, 60А" in docx.text
    xlsx = extract_text("benefits.xlsx", None, _xlsx_bytes())
    assert "# Лист: Льготы" in xlsx.text and "ДМС | после испытательного срока" in xlsx.text
    pdf = extract_text("kb.pdf", "application/pdf", _pdf_bytes())
    assert "Napoleon IT knowledge base" in pdf.text and pdf.content_type == "application/pdf"
    # Формат по content-type, если расширения нет.
    assert extract_text(None, "text/markdown", "# Заголовок".encode()).extension == ".md"


def test_extract_text_rejects_unknown_empty_and_oversized() -> None:
    with pytest.raises(UnsupportedFormatError):
        extract_text("archive.zip", "application/zip", b"PK")
    with pytest.raises(UnsupportedFormatError):
        extract_text(None, None, b"data")
    with pytest.raises(ParseError):
        extract_text("empty.txt", None, b"   \n  ")
    with pytest.raises(ParseError):
        extract_text("broken.docx", None, b"not a docx")
    with pytest.raises(Exception, match="больше"):
        extract_text("big.txt", None, b"x" * (parsing.MAX_FILE_BYTES + 1))


def test_chunk_text_splits_paragraphs_and_merges_tail() -> None:
    long_paragraph = " ".join(f"Предложение номер {i} про стек компании." for i in range(120))
    chunks = chunk_text(f"Короткий абзац.\n\n{long_paragraph}\n\nХвост.", size=600)
    assert len(chunks) >= 3
    assert all(len(chunk) <= 600 for chunk in chunks)
    assert chunks[0].startswith("Короткий абзац.")
    assert chunks[-1].endswith("Хвост.")
    # Соседние короткие абзацы склеиваются в один фрагмент, длинный не ломается.
    short_pair = "А" * 500 + "\n\n" + "Б" * 50
    assert chunk_text(short_pair, size=600) == [short_pair]
    assert [len(c) for c in chunk_text("В" * 1500, size=600)] == [600, 600, 300]
    assert chunk_text("") == []


# ------------------------------------------------------------------- поиск


def test_rank_uses_light_russian_stemming_and_ignores_stopwords() -> None:
    passages = [
        search.Passage("hiring", "Найм", "Этапы отбора: отклик, интервью с руководителем, оффер."),
        search.Passage("intern", "Стажировка", "NapTalentBootcamp — стажировка для студентов."),
        search.Passage("stack", "Стек", "Бэкенд: Python, Java, PHP. Фронтенд: Vue.js, React."),
    ]
    assert search.rank("стажировки для студента", passages)[0].key == "intern"
    assert search.rank("какой у вас стек бэкенда", passages)[0].key == "stack"
    assert search.rank("и в на", passages) == []
    assert search.rank("", passages) == []
    assert search.stem("стажировки") == search.stem("стажировка")
    assert search.tokenize("Ёлки и python-разработчики") == ["елки", "python-разработчик"]


# --------------------------------------------------------------------- API


async def test_documents_crud_search_and_roles(client: AsyncClient) -> None:
    _, owner = await register(client)
    created = await client.post(
        "/api/knowledge",
        json={"title": "О компании", "text": COMPANY_TEXT, "tags": [" Компания ", "hr", "hr"]},
        headers=bearer(owner),
    )
    assert created.status_code == 201, created.text
    document = created.json()
    assert document["kind"] == "text" and document["chunk_count"] >= 1
    assert document["tags"] == ["компания", "hr"]
    assert document["preview"].startswith("# О компании")

    listed = (await client.get("/api/knowledge", headers=bearer(owner))).json()
    assert [item["id"] for item in listed] == [document["id"]]

    found = await client.get(
        "/api/knowledge/search", params={"q": "этапы найма"}, headers=bearer(owner)
    )
    assert found.status_code == 200
    body = found.json()
    assert body["documents_total"] == 1 and body["hits"]
    assert body["hits"][0]["document_id"] == document["id"]
    assert "отбора" in body["hits"][0]["text"]

    # Нанимающий менеджер читает и ищет, но не пополняет базу.
    manager = await _manager(client, owner)
    assert (await client.get("/api/knowledge", headers=bearer(manager))).status_code == 200
    denied = await client.post(
        "/api/knowledge", json={"title": "x", "text": "y" * 30}, headers=bearer(manager)
    )
    assert denied.status_code == 403
    assert (
        await client.delete(f"/api/knowledge/{document['id']}", headers=bearer(manager))
    ).status_code == 403

    # Правка текста перенарезает фрагменты.
    updated = await client.patch(
        f"/api/knowledge/{document['id']}",
        json={"title": "О компании и найме", "text": "Только одна строка."},
        headers=bearer(owner),
    )
    assert updated.status_code == 200
    assert updated.json()["title"] == "О компании и найме"
    assert updated.json()["chunk_count"] == 1 and updated.json()["text"] == "Только одна строка."

    # Чужая организация документ не видит.
    _, stranger = await register(client)
    assert (
        await client.get(f"/api/knowledge/{document['id']}", headers=bearer(stranger))
    ).status_code == 404

    deleted = await client.delete(f"/api/knowledge/{document['id']}", headers=bearer(owner))
    assert deleted.status_code == 204
    assert (
        await client.get(f"/api/knowledge/{document['id']}", headers=bearer(owner))
    ).status_code == 404
    formats = (await client.get("/api/knowledge/formats", headers=bearer(owner))).json()
    assert ".pdf" in formats["extensions"] and formats["max_file_mb"] == 20


async def test_upload_parses_file_and_stores_original(client: AsyncClient) -> None:
    from leonit.core.storage import get_storage

    _, owner = await register(client)
    response = await client.post(
        "/api/knowledge/upload",
        files={"file": ("Ценности компании.docx", _docx_bytes(), "application/octet-stream")},
        data={"tags": "культура, hr"},
        headers=bearer(owner),
    )
    assert response.status_code == 201, response.text
    document = response.json()
    assert document["kind"] == "file" and document["title"] == "Ценности компании"
    assert document["source_name"] == "Ценности компании.docx"
    assert document["tags"] == ["культура", "hr"] and document["size_bytes"] > 0
    detail = (await client.get(f"/api/knowledge/{document['id']}", headers=bearer(owner))).json()
    assert "Learn IT" in detail["text"]
    hit = (
        await client.get("/api/knowledge/search", params={"q": "где офис"}, headers=bearer(owner))
    ).json()["hits"][0]
    assert "Челябинск" in hit["text"]

    storage = get_storage()
    prefix = "knowledge/"
    # Файл лежит в хранилище и исчезает вместе с документом.
    from sqlalchemy import select

    from leonit.core.db import get_session_maker
    from leonit.knowledge.models import KnowledgeDocument

    async with get_session_maker()() as session:
        key = await session.scalar(
            select(KnowledgeDocument.storage_key).where(
                KnowledgeDocument.id == uuid.UUID(document["id"])
            )
        )
    assert key and key.startswith(prefix) and await storage.exists(key)
    assert (
        await client.delete(f"/api/knowledge/{document['id']}", headers=bearer(owner))
    ).status_code == 204
    assert not await storage.exists(key)

    rejected = await client.post(
        "/api/knowledge/upload",
        files={"file": ("photo.png", b"\x89PNG", "image/png")},
        headers=bearer(owner),
    )
    assert rejected.status_code == 422 and "не поддерживается" in rejected.json()["detail"]


# --------------------------------------------------------------- ассистент


async def test_assistant_searches_knowledge_base(client: AsyncClient) -> None:
    _, owner = await register(client)
    thread = await _thread(client, owner, page_path="/organization/knowledge")
    empty = await _send(
        client, owner, thread["id"], '[[call:search_knowledge {"query": "стажировка"}]]'
    )
    action = _actions(empty)[0]
    assert action["tool"] == "search_knowledge" and action["kind"] == "done"
    assert "пуста" in action["summary"] and action["result"] == []

    await client.post(
        "/api/knowledge",
        json={"title": "О компании", "text": COMPANY_TEXT},
        headers=bearer(owner),
    )
    found = await _send(
        client, owner, thread["id"], '[[call:search_knowledge {"query": "стажировка"}]]'
    )
    action = _actions(found)[0]
    assert action["kind"] == "done" and action["result"]
    assert action["result"][0]["title"] == "О компании"
    assert "NapTalentBootcamp" in action["result"][0]["text"]


# ------------------------------------------------- быстрый черновик вакансии


VACANCY_TEXT = """Python-разработчик (middle), Челябинск или удалённо.

Обязанности: разработка бэкенда на FastAPI, интеграции с внешними API,
оптимизация PostgreSQL. Требования: Python от 3 лет, asyncio, Docker,
опыт с очередями задач. Условия: ДМС, гибкий график, обучение за счёт компании.
"""


async def test_quick_vacancy_from_text_builds_rubric_and_questions(client: AsyncClient) -> None:
    _, owner = await register(client)
    await client.post(
        "/api/knowledge",
        json={"title": "О компании", "text": COMPANY_TEXT},
        headers=bearer(owner),
    )
    response = await client.post(
        "/api/vacancies/quick", json={"text": VACANCY_TEXT}, headers=bearer(owner)
    )
    assert response.status_code == 201, response.text
    body = response.json()
    vacancy = body["vacancy"]
    assert vacancy["status"] == "draft" and vacancy["title"]
    assert vacancy["rubric"] and all(item["id"] for item in vacancy["rubric"])
    assert len({item["id"] for item in vacancy["rubric"]}) == len(vacancy["rubric"])
    assert vacancy["questions"] and vacancy["question_count"] == len(vacancy["questions"])
    known = {item["id"] for item in vacancy["rubric"]}
    assert all(set(q["competency_ids"]) <= known for q in vacancy["questions"])
    assert isinstance(body["notes"], str)
    # Черновик виден в списке и редактируется как обычная вакансия.
    listed = (await client.get("/api/vacancies", headers=bearer(owner))).json()
    assert [item["id"] for item in listed] == [vacancy["id"]]

    short = await client.post(
        "/api/vacancies/quick", json={"text": "Python"}, headers=bearer(owner)
    )
    assert short.status_code == 422

    manager = await _manager(client, owner)
    forbidden = await client.post(
        "/api/vacancies/quick", json={"text": VACANCY_TEXT}, headers=bearer(manager)
    )
    assert forbidden.status_code == 403


async def test_quick_vacancy_from_uploaded_file(client: AsyncClient) -> None:
    _, owner = await register(client)
    response = await client.post(
        "/api/vacancies/quick/upload",
        files={"file": ("vacancy.md", VACANCY_TEXT.encode("utf-8"), "text/markdown")},
        headers=bearer(owner),
    )
    assert response.status_code == 201, response.text
    body = response.json()
    assert body["source_name"] == "vacancy.md"
    assert body["vacancy"]["rubric"] and body["vacancy"]["questions"]

    rejected = await client.post(
        "/api/vacancies/quick/upload",
        files={"file": ("vacancy.exe", b"MZ", "application/octet-stream")},
        headers=bearer(owner),
    )
    assert rejected.status_code == 422
